#!/usr/bin/env python3
"""W01 javob varaqasini to'ldiradi.

Ishlatish:
  python3 w01_fill.py BO'SH.docx NATIJA.docx [amaliy.json]

amaliy.json — amaliy qismlar uchun (srv1 da olingan haqiqiy natijalar):
  {"text": {"L1-1": "...", ...}, "image": {"L1-3": "/yo'l/L1-3.png", ...}}
"""
import json
import sys

import docx
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

HEADER = {
    "Familiya va ism": "Erdonov Ravshan",
    "Talaba ID": "",
    "Guruh": "201 IT National",
    "Topshirilgan sana": "2026-10-23",
}

MC = {
    "M1.1": "B", "M1.2": "C", "M1.3": "A", "M1.4": "C", "M1.5": "B",
    "M2.1": "A", "M2.2": "D", "M2.3": "C", "M2.4": "A",
    "M3.1": "B", "M3.2": "C", "M3.3": "B", "M3.4": "B", "M3.5": "D",
    "M4.1": "B", "M4.2": "D",
}

TEXT = {
    "M1.6": "Birinchi hodisa faqat bitta dastur (veb-server) yiqilganini aytadi - bu user space dagi "
            "dastur, mashina undan keyin ham ishlashi kerak edi. Butun mashina to'xtagani esa kernel "
            "ishdan chiqqanini bildiradi, chunki tizimni faqat kernel ning ishdan chiqishi to'liq to'xtatadi.",
    "M1.7": "Linux - bu faqat kernel, uni 1991-yilda Linus Torvalds e'lon qilgan. GNU/Linux esa to'liq "
            "operatsion tizim: Linux kernel + GNU loyihasining vositalari (shell, kompilyator, asosiy "
            "buyruqlar). Kundalik ishlatadigan buyruqlarimizning ko'pi aslida GNU dan keladi.",
    "M2.6": "Ziddiyat emas. NAT rejimida har bir guest ga hypervisor o'zining alohida yopiq tarmog'ini "
            "beradi, shuning uchun ikkala mashina ham 10.0.2.15 olishi mumkin va ular bir-birini "
            "ko'rmaydi. Xuddi har xil binolarda 3-kvartira bo'lgani kabi.",
    "M2.7": "1) Bridged da manzilni universitet tarmog'i beradi, biz uni boshqarmaymiz - manzil o'zgarib "
            "qolishi mumkin.\n2) Universitet tarmog'i bunday qurilmalarni cheklashi mumkin, simsiz "
            "tarmoqlar esa bridged ulanishni ko'pincha umuman qabul qilmaydi.",
    "M2.8": "Hech narsa yo'qolmagan. Disk dinamik ajratilgan: 25 GB - bu chegara, hozir ajratilgan joy "
            "emas. Fayl guest yozgan hajmicha (6 GB) bo'ladi va kerak bo'lsa 25 GB gacha o'sadi.",
    "M3.7": "Yo'qoladi: snapshot dan keyin (shu bir oy ichida) mashinada yaratilgan fayllar va o'rnatilgan "
            "dasturlar, masalan oxirgi oydagi yangilanishlar.\nSaqlanadi: mashinadan tashqarida "
            "turgan narsalar - host dagi server kartasi va qaydnoma, .ova eksport fayli.",
    "M4.3": "Yangi dastur taxminan 1.5Gi ola oladi - bu available ustuni. free (1.3Gi) faqat umuman "
            "tegilmagan xotira. buff/cache dagi 404Mi ni kernel disk keshi uchun ushlab turibdi, lekin "
            "dasturga kerak bo'lsa darhol qaytarib beradi, shuning uchun available free dan katta.",
    "M4.4": "Nosozlik emas. Guided LVM tuzilmasi disk joyining taxminan yarmini / ga beradi, qolganini "
            "keyin kengaytirish uchun bo'sh qoldiradi. /dev/mapper/ubuntu--vg-ubuntu--lv degan nom / "
            "LVM ning mantiqiy hajmida (ubuntu-vg guruhidagi ubuntu-lv) turganini ko'rsatadi.",
    "M4.5": "Avval qaysi 3 ta package olib tashlanishini o'qib chiqaman va ular nima ekanini tekshiraman. "
            "Oddiy yangilanishda package o'chirilishi g'alati holat - kerakli xizmat yoki kernel o'chib "
            "ketishi mumkin. Tushunmasam, n deb javob beraman va avval snapshot olaman.",
    "M4.6": "Oynani yopib Power off qilish - quvvat simini sug'urish bilan bir xil: diskka yozilayotgan "
            "ma'lumot yo'qolishi, keyingi yuklanishda fayl tizimi tiklash talab qilishi mumkin. To'g'ri "
            "usul - sudo poweroff: xizmatlarni tartib bilan to'xtatadi, yozuvni yakunlab, o'chiradi.",
    "M5.1": "2026-10-23  sudo apt upgrade  srv1  mashina ortda qolgan edi  ikkinchi sudo apt update "
            "yangilanadigan narsa qolmaganini ko'rsatdi, mashina qayta yuklandi va sysadmin kira oldi",
    "M5.2": "(b) firewall ni o'zgartirish - tizimga kirish imkonini yo'qotishi mumkin.\n(c) /srv/company "
            "ostidagi fayllarni o'chirish - ma'lumotni yo'q qiladi.\nIkkalasidan oldin snapshot "
            "olinadi. (a) /etc/os-release ni o'qish hech narsani o'zgartirmaydi, snapshot kerak emas.",
    "M5.3": "1) Kartaga parol yozilgan - karta qo'ldan qo'lga o'tadi, parol yozilmasligi kerak (faqat "
            "u qayerda saqlanishi yoziladi).\n2) Parolning o'zi juda oddiy (Pa55word) - taxmin qilish "
            "oson. Uni almashtirish kerak, chunki u endi ochiq bo'lib qolgan.",
    "M5.4": "Log fayllarda srv1 ko'rinadi - bu hostname, mashina ichidagi nom. srv1-karimov faqat "
            "VirtualBox ro'yxatidagi yorliq: bitta kompyuterda bir nechta talabaning mashinasini "
            "farqlash uchun. Mashinaning o'zi bu yorliqni bilmaydi.",
    "L2-1": "Hozir fayllar xodimlarning noutbuklarida tarqoq yotibdi va flesh orqali yuriladi - ikkita "
            "fayl allaqachon yo'qolgan. Tashkilotga fayllar bitta joyda, bo'limlar bo'yicha ajratilgan "
            "va zaxiralangan holda saqlanishi kerak. Bundan tashqari, brauzerda ishlaydigan bitta ichki "
            "dastur va tashqaridan xavfsiz ulanish kerak.",
    "X1": "Ubuntu Server hujjatlari:\n- Bilmagan 3 ta narsa: o'rnatishda tarmoq manzilini qo'lda (static) "
          "kiritsa bo'ladi; SSH kalitni o'rnatish vaqtida GitHub/Launchpad dan import qilish mumkin; "
          "birinchi yuklanishda sozlamalarni cloud-init bajaradi.\n- Tushunmaganim: autoinstall fayli "
          "qanday yoziladi.\nVirtualBox qo'llanmasi (First Steps):\n- Bilmagan 3 ta narsa: diskning VDI "
          "dan boshqa formatlari ham bor (VMDK, VHD); mashinani Clone qilish mumkin; Guest Additions "
          "umumiy papka va clipboard beradi.\n- Tushunmaganim: linked clone oddiy clone dan nimasi bilan "
          "farq qiladi.",
    "X2": "Minimized variant odam kirib ishlashi uchun kerak bo'lgan ko'p narsalarni olib tashlaydi: "
          "man sahifalar, muharrirlar va ko'plab standart vositalar. U avtomatik boshqariladigan, "
          "odam kirmaydigan serverlar uchun. Biz serverga kirib o'rganamiz, keyin bu vositalar kerak "
          "bo'ladi - shuning uchun to'liq Ubuntu Server tanlandi.",
    "X3": "Guided LVM diskda /boot uchun alohida bo'lim va qolgan joyda LVM volume group (ubuntu-vg) "
          "yaratadi. Uning ichida / uchun logical volume (ubuntu-lv) ochiladi, guruhdagi joyning bir "
          "qismi keyin kengaytirish uchun bo'sh qoladi.",
    "X4": "Internal network - faqat shu tarmoqqa ulangan guest lar bir-birini ko'radi, host ham, "
          "internet ham yo'q. Bizga host dan serverga kirish kerak, shuning uchun Host-only olindi. "
          "NAT network - bir nechta guest bitta NAT ortida bir-birini ko'rib, internetga chiqadi. "
          "Bizda internet uchun oddiy NAT yetarli, ichki aloqa Host-only orqali.",
    "X5": "Qo'llanma guest ga host xotirasining katta qismini bermaslikni aytadi: guest ishlayotganda "
          "o'sha xotira host dan olinadi. Ortiqcha bersangiz, host ning o'zi sekinlashadi yoki "
          "swap ga tushib, ikkalasi ham qotib qoladi.",
    "P1": "Mashq mashinasi yo'q, tartibni xotiradan yozdim:\n1) VirtualBox - New, nom practice, "
          "Linux/Ubuntu 64-bit, ISO tanlanadi, unattended o'chiriladi.\n2) 1 vCPU, 1024 MB, 10 GB "
          "dynamic.\n3) Settings - Network: 1-adapter NAT, 2-adapter Host-only.\n4) Start - English, "
          "klaviatura, Ubuntu Server (minimized emas), tarmoq/proxy/mirror o'zgarmaydi, Use entire "
          "disk + LVM, Continue, profil: practice / sysadmin, Pro - Skip, OpenSSH belgilanadi, snaps yo'q.\n"
          "5) Reboot Now, keyin ISO uziladi, Enter.\n6) Kirish: uname -a, cat /etc/os-release, free -h, "
          "df -h.\nAdashgan joylarim: ISO ni Reboot Now dan oldin uzish kerak deb o'ylagandim (aslida "
          "keyin); OpenSSH standart belgilangan deb o'ylagandim (aslida o'zim Space bosishim kerak).",
}

M36 = [
    ["Mashinaning shu paytdagi to'liq holati (disk, sozlamalar)",
     "Butun mashina, bitta faylga", "Kerakli ma'lumot (fayllar)"],
    ["Mashina bilan bir katalogda, o'sha diskda",
     "Istalgan joyda, boshqa diskka ham ko'chirsa bo'ladi", "Boshqa saqlash qurilmasida"],
    ["Hech narsa - mashina bilan birga yo'qoladi",
     "Boshqa diskda bo'lsa saqlanadi (eksport kunidagi holat)", "Saqlanadi"],
    ["Yo'q, faqat butun holat qaytadi",
     "Yo'q, faqat butun mashinani import qilish mumkin", "Ha"],
]

GLOSSARY = {
    "kernel": "OS ning yuragi: protsessor vaqti, xotira, qurilmalar, fayllar va process larni boshqaradi.",
    "user space": "kernel dan tashqaridagi hamma dasturlar - shell, buyruqlar, xizmatlar.",
    "system call": "Dastur kernel dan biror resurs so'raydigan yagona yo'l.",
    "ochiq kod": "Dasturni ishlatish, o'qish, o'zgartirish va tarqatish huquqi.",
    "GPL": "Linux kernel litsenziyasi: to'rt erkinlik + copyleft sharti.",
    "copyleft": "O'zgartirib tarqatsang, o'sha huquqlarni va manba kodini ham berishing shart.",
    "distribution": "Kernel, vositalar va dastur o'rnatish usuli yig'ilgan tayyor Linux (masalan Ubuntu).",
    "hypervisor": "Virtual mashinalarni yaratib, bir-biridan ajratib turadigan dastur (VirtualBox).",
    "host": "VM ishlab turgan haqiqiy kompyuter va uning OS i.",
    "guest": "VM ichida ishlayotgan OS.",
    "1-tur va 2-tur": "1-tur temirning ustida o'zi ishlaydi (ESXi, KVM); 2-tur oddiy OS ichida dastur (VirtualBox).",
    "vCPU": "guest ga berilgan virtual protsessor yadrosi.",
    "virtual disk": "host dagi bitta fayl, guest uni to'liq disk deb ko'radi.",
    "dinamik ajratilgan": "Disk fayli yozilgan sari o'sadi, e'lon qilingan hajm - faqat chegara.",
    "ISO obraz": "O'rnatish diskining bitta fayldagi nusxasi.",
    "live installer": "ISO dan xotirada ishga tushib, tizimni diskka o'rnatadigan kichik Linux.",
    "LVM": "Disk bilan fayl tizimi orasidagi qatlam; keyin / ni kengaytirish imkonini beradi.",
    "hostname": "Mashinaning o'z nomi - prompt da, log larda va tarmoqda ko'rinadi.",
    "sudo": "Bitta buyruqqa admin huquqi, o'z parolim bilan.",
    "NAT": "guest internetga host orqali chiqadi, tashqaridan unga kirib bo'lmaydi.",
    "Host-only": "Kompyuter ichidagi yopiq tarmoq: host va guest lar bir-birini ko'radi.",
    "Bridged": "guest haqiqiy tarmoqqa alohida qurilma bo'lib chiqadi; bizda ishlatilmaydi.",
    "repository": "Distribution tekshirib, imzolab beradigan package lar ombori.",
    "package ro'yxati": "Repository dagi package larning mashinadagi nusxasi; apt update yangilaydi.",
    "snapshot": "VM ning ma'lum paytdagi to'liq holati, mashina bilan bir joyda saqlanadi.",
    "appliance export": "Butun mashinani boshqa joyga olib borsa bo'ladigan bitta .ova fayl.",
    "server kartasi": "Mashina haqida bir sahifa: mashinadan o'qib yoziladi, parolsiz.",
    "o'zgarishlar qaydnomasi": "Har bir o'zgarishga sanali bitta satr: nima, qayerda, nega, qanday tekshirildi.",
}


def write(cell, text):
    """Katakdagi bo'sh paragrafga matn yozadi (qatorlar \\n bilan)."""
    p = cell.paragraphs[0]
    for i, line in enumerate(text.split("\n")):
        r = p.add_run(line)
        r.font.name = "Inter"
        r.font.size = Pt(10)
        if i < len(text.split("\n")) - 1:
            r.add_break()


def label(cell):
    return cell.text.split("\n")[0].strip()


def main():
    src, dst = sys.argv[1], sys.argv[2]
    extra = {"text": {}, "image": {}}
    if len(sys.argv) > 3:
        with open(sys.argv[3], encoding="utf-8") as f:
            extra.update(json.load(f))
    text = dict(TEXT, **extra.get("text", {}))
    images = extra.get("image", {})

    d = docx.Document(src)
    for t in d.tables:
        first = label(t.rows[0].cells[0])
        if first == "Familiya va ism":
            for r in t.rows:
                v = HEADER.get(label(r.cells[0]))
                if v:
                    write(r.cells[1], v)
        elif first == "Mezon":
            for r, vals in zip(t.rows[1:], M36):
                for c, v in zip(r.cells[1:], vals):
                    write(c, v)
        elif first == "Atama":
            for r in t.rows[1:]:
                v = GLOSSARY.get(label(r.cells[0]))
                if v:
                    write(r.cells[1], v)
        else:
            cells = t.rows[0].cells
            for row in t.rows:
                cells = row.cells
                for i, c in enumerate(cells[:-1]):
                    qid = label(c)
                    if qid in MC:
                        sdt = cells[i + 1]._tc.find(".//" + qn("w:sdt"))
                        ph = sdt.find(qn("w:sdtPr") + "/" + qn("w:showingPlcHdr"))
                        if ph is not None:
                            ph.getparent().remove(ph)
                        for rpr in sdt.iter(qn("w:color")):
                            rpr.set(qn("w:val"), "0A1A2E")
                        sdt.find(".//" + qn("w:t")).text = MC[qid]
                    elif qid in text:
                        write(cells[i + 1], text[qid])
                    elif qid in images:
                        target = cells[i + 1]
                        for p in target.paragraphs:
                            for r in p.runs:
                                if "rasm shu yerga" in r.text:
                                    r.text = ""
                        target.paragraphs[0].add_run().add_picture(images[qid], width=Cm(14))
    d.save(dst)
    print("saqlandi:", dst)


if __name__ == "__main__":
    main()
